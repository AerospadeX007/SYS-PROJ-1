"""System-Level Results 1 (S.2.3) as native, editable charts.

1. Excel: a workbook with the ten team designs (vehicle level vs system level) and a native Excel
   scatter chart of gross liftoff mass vs program cost, one line per design from its Section 1
   (vehicle-level) point to its system-level point.
2. PowerPoint: the same chart as a native PowerPoint chart (right-click > Edit Data opens it in
   Excel), replacing the picture on the "System-Level Results 1" slide of the group deck.

Inputs come from main_section2.py's output, csvs/S2_report.json (S2_1.py results at each teammate's
Section 1 min-mass and min-cost splits).

    python S2_results_chart.py                     -> csvs/S2_system_level_results.xlsx
    python S2_results_chart.py <report.json> <out.xlsx> [<deck.pptx> <out.pptx>]
    (run python main_section2.py first; it writes csvs/S2_report.json)

Needs openpyxl and python-pptx (pip install openpyxl python-pptx).
Author: Shaunn Pavelik (10/15/2026)
"""
import copy
import json
import sys

from lxml import etree
from openpyxl import Workbook
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.chart.legend import Legend, LegendEntry
from openpyxl.chart.marker import DataPoint, Marker
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import SeriesLabel
from openpyxl.chart.data_source import StrRef
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText
from openpyxl.drawing.line import LineEndProperties, LineProperties
from openpyxl.drawing.text import CharacterProperties, Paragraph, ParagraphProperties, Font as DFont
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------- data
TEAM = ["Gregory Kahn", "Henry Rowley", "Jacob Harmon", "Sharan Menon", "Shaunn Pavelik"]
COLOR = {"Gregory Kahn": "2A78D6", "Henry Rowley": "1BAF7A", "Jacob Harmon": "EB6834",
         "Sharan Menon": "8E5BB5", "Shaunn Pavelik": "E0A526"}          # same as S2_auto_graph
VL_GRAY, INK = "A5A5A0", "1A1A1A"
PICKS = {("Henry Rowley", "min-mass"), ("Shaunn Pavelik", "min-cost")}  # group selections (S.2.3)
XLIM, YLIM = (1200, 2250), (9.5, 12.5)                                  # chart window (t, $B)
TINT = {"Gregory Kahn": "B9D3F2", "Henry Rowley": "B3E6D3", "Jacob Harmon": "F7C3AA",
        "Sharan Menon": "D8C6E6", "Shaunn Pavelik": "F4DDA1"}           # light fill = engines don't fit


def pick_label(r):
    return f"{'Lowest mass' if r['kind'] == 'min-mass' else 'Lowest cost'}: {r['sl_m']:,.0f} t, ${r['sl_c']:.2f}B"


def legend_order(rows, who):
    """this teammate's on-chart designs, non-picks first so the legend entry is the pair name"""
    mine = [r for r in rows if r["who"] == who and on_chart(r)]
    return sorted(mine, key=lambda r: r["pick"])
DISPLAY = {"SOLID": "Solid"}


def load(report_path):
    people = json.load(open(report_path))["people"]
    rows = []
    for who in TEAM:
        e = people[who]
        for kind in ("min-mass", "min-cost"):
            d = e["designs"][kind]
            rows.append(dict(who=who, last=who.split()[-1], s1=DISPLAY.get(e["pair"][0], e["pair"][0]), s2=e["pair"][1],
                             kind=kind, dv1=d["dv1"], vl_m=d["vl_m0"] / 1e3, sl_m=d["sl_m0"] / 1e3,
                             vl_c=d["vl_cost"] / 1e3, sl_c=d["sl_cost"] / 1e3, D=d["d"], n1=d["n1"], n2=d["n2"],
                             fit=bool(d["fit"]), pick=(who, kind) in PICKS))
    return rows


def on_chart(r):
    return XLIM[0] <= r["sl_m"] <= XLIM[1] and YLIM[0] <= r["sl_c"] <= YLIM[1]


def off_chart_note(rows):
    off = [f"{r['s1']} / {r['s2']} {r['kind']}: {r['sl_m']:,.0f} t, ${r['sl_c']:.2f}B" for r in rows if not on_chart(r)]
    return "Off the chart: " + "; ".join(off) if off else ""


# ---------------------------------------------------------------- Excel
def build_workbook(rows, out):
    wb = Workbook()
    ws = wb.active
    ws.title = "S2.3 Results"
    arial = lambda **k: Font(name="Arial", **k)
    thin = Side(style="thin", color="B8BFCC")
    head_fill = PatternFill("solid", fgColor="011893")

    ws["A1"] = "Team 3 – System-Level Results 1 (S.2.3): vehicle level vs system level"
    ws["A1"].font = arial(bold=True, size=14, color="011893")
    ws["A2"] = ("Source: main_section2.py (S2_1.py), csvs/S2_report.json. VL = Section 1 vehicle level (Table 2 δ and Isp); "
                "SL = system level (Lecture 8 MERs + 30% margin, J-2X upper stage). Blue = model output, black = formula.")
    ws["A2"].font = arial(italic=True, size=9, color="4A5568")

    heads = ["Teammate", "Stage 1", "Stage 2", "Design", "ΔV₁ (m/s)", "VL gross mass (t)", "SL gross mass (t)",
             "Mass change", "VL cost ($B FY2025)", "SL cost ($B FY2025)", "Cost change", "Diameter (m)",
             "Engines S1", "Engines S2", "Engines fit", "Group pick"]
    widths = [11, 12, 10, 10, 10, 13, 13, 11, 13, 13, 11, 11, 10, 10, 10, 11]
    R0 = 4
    for j, (h, w) in enumerate(zip(heads, widths), 1):
        c = ws.cell(R0, j, h)
        c.font = arial(bold=True, color="FFFFFF", size=10)
        c.fill = head_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[R0].height = 30

    blue = "0000FF"
    for i, r in enumerate(rows):
        n = R0 + 1 + i
        vals = [r["last"], r["s1"], r["s2"], "Min mass" if r["kind"] == "min-mass" else "Min cost", r["dv1"],
                r["vl_m"], r["sl_m"], f"=G{n}/F{n}-1", r["vl_c"], r["sl_c"], f"=J{n}/I{n}-1", r["D"],
                r["n1"], r["n2"], "yes" if r["fit"] else "no", "yes" if r["pick"] else ""]
        fmts = [None, None, None, None, "#,##0", "#,##0", "#,##0", "+0.0%;-0.0%", "$0.00", "$0.00", "+0.0%;-0.0%",
                "0.0", "0", "0", None, None]
        for j, (v, f) in enumerate(zip(vals, fmts), 1):
            c = ws.cell(n, j, v)
            is_formula = isinstance(v, str) and v.startswith("=")
            c.font = arial(size=10, bold=r["pick"], color=("000000" if is_formula or j <= 4 else blue))
            if j == 1:
                c.fill = PatternFill("solid", fgColor={"Gregory Kahn": "FFF1C9", "Henry Rowley": "DCEBFA",
                                                       "Jacob Harmon": "FADCDC", "Sharan Menon": "DDF3E4",
                                                       "Shaunn Pavelik": "E9DDF7"}[r["who"]])
            if f:
                c.number_format = f
            c.alignment = Alignment(horizontal="left" if j <= 4 else "center")
            c.border = Border(bottom=thin)
            if j == 15 and not r["fit"]:
                c.font = arial(size=10, bold=True, color="B4231F")
    last = R0 + len(rows)

    # helper ranges the chart reads (formulas into the table, so the chart follows edits)
    H0 = last + 3
    ws.cell(H0, 1, "Chart helper: group picks (labeled on the chart)").font = arial(bold=True, size=10)
    ws.cell(H0 + 1, 1, "Design").font = arial(bold=True, size=10)
    ws.cell(H0 + 1, 2, "SL gross mass (t)").font = arial(bold=True, size=10)
    ws.cell(H0 + 1, 3, "SL cost ($B)").font = arial(bold=True, size=10)
    ws.cell(H0 + 1, 4, "Chart label").font = arial(bold=True, size=10)
    picks_rows = [R0 + 1 + i for i, r in enumerate(rows) if r["pick"]]
    for k, n in enumerate(picks_rows):
        h = H0 + 2 + k
        ws.cell(h, 1, f'=A{n}&" "&D{n}').font = arial(size=10)
        ws.cell(h, 2, f"=G{n}").number_format = "#,##0"
        ws.cell(h, 3, f"=J{n}").number_format = "$0.00"
        kind = "Lowest mass" if rows[n - R0 - 1]["kind"] == "min-mass" else "Lowest cost"
        ws.cell(h, 4, f'="{kind}: "&TEXT(G{n},"#,##0")&" t, $"&TEXT(J{n},"0.00")&"B"').font = arial(size=10)
        for j in (2, 3):
            ws.cell(h, j).font = arial(size=10)
    note = off_chart_note(rows)
    ws.cell(H0 + 5, 1, note + ". Circle = min-mass, square = min-cost, light fill = engines don't fit under the stage.").font = \
        arial(italic=True, size=9, color="4A5568")

    # ---- native Excel scatter chart
    ch = ScatterChart(scatterStyle="lineMarker")
    ch.title = "Vehicle level (gray) to system level (color); labeled = group picks"
    ch.style = 2
    ch.height, ch.width = 13, 24
    ch.x_axis.title = "Gross liftoff mass (t)"
    ch.y_axis.title = "NRE program cost ($B, FY2025)"
    ch.x_axis.scaling.min, ch.x_axis.scaling.max = XLIM
    ch.y_axis.scaling.min, ch.y_axis.scaling.max = YLIM
    ch.x_axis.majorUnit, ch.y_axis.majorUnit = 200, 0.5
    ch.x_axis.number_format, ch.y_axis.number_format = "#,##0", "0.0"
    ch.x_axis.delete = ch.y_axis.delete = False
    ch.legend = Legend()
    ch.legend.position = "r"
    hidden = []

    def add(series):
        ch.series.append(series)
        return len(ch.series) - 1

    def gp(series):
        """series line properties, created if openpyxl left them empty"""
        if series.graphicalProperties is None:
            series.graphicalProperties = GraphicalProperties()
        if series.graphicalProperties.line is None:
            series.graphicalProperties.line = LineProperties()
        return series.graphicalProperties

    # 2) gray vehicle-level points
    s = Series(Reference(ws, min_col=9, min_row=R0 + 1, max_row=last), Reference(ws, min_col=6, min_row=R0 + 1, max_row=last),
               title="Vehicle level (Section 1)")
    s.marker = Marker(symbol="circle", size=6)
    s.marker.graphicalProperties = GraphicalProperties(solidFill=VL_GRAY)
    s.marker.graphicalProperties.line.solidFill = VL_GRAY
    gp(s).line.noFill = True
    add(s)
    # 3) one arrow per design, Section 1 point -> system-level point (hidden from the legend)
    for i, r in enumerate(rows):
        if not on_chart(r):
            continue
        n = R0 + 1 + i
        s = Series(Reference(ws, min_col=9, max_col=10, min_row=n), Reference(ws, min_col=6, max_col=7, min_row=n),
                   title=f"{r['last']} {r['kind']} (line)")
        s.marker = Marker(symbol="none")
        gp(s).line.solidFill = COLOR[r["who"]]
        gp(s).line.width = 19050
        gp(s).line.tailEnd = LineEndProperties(type="triangle", w="med", len="med")
        hidden.append(add(s))
    # 4) system-level point per design: circle = min-mass, square = min-cost, light fill = engines don't
    #    fit. One legend entry per teammate; the two group picks are bigger and labeled.
    pick_cells = {n: H0 + 2 + k for k, n in enumerate(picks_rows)}
    for who in TEAM:
        for j, r in enumerate(legend_order(rows, who)):
            n = R0 + 1 + rows.index(r)
            title = f"{r['s1']} / {r['s2']} ({r['last']})"
            if r["pick"]:
                title = SeriesLabel(strRef=StrRef(f"'{ws.title}'!$D${pick_cells[n]}"))
            s = Series(Reference(ws, min_col=10, min_row=n), Reference(ws, min_col=7, min_row=n), title=title)
            if r["pick"]:
                s.tx = title
            s.marker = Marker(symbol="circle" if r["kind"] == "min-mass" else "square", size=14 if r["pick"] else 10)
            s.marker.graphicalProperties = GraphicalProperties(solidFill=COLOR[who] if r["fit"] else TINT[who])
            s.marker.graphicalProperties.line.solidFill = INK if r["pick"] else COLOR[who]
            s.marker.graphicalProperties.line.width = 25400
            gp(s).line.noFill = True
            if r["pick"]:
                s.dLbls = DataLabelList()
                s.dLbls.showSerName = True
                s.dLbls.showVal = s.dLbls.showCatName = s.dLbls.showLegendKey = False
                s.dLbls.position = "b"
                s.dLbls.txPr = RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=CharacterProperties(b=True, sz=900)),
                                                     endParaRPr=CharacterProperties())])
            k = add(s)
            if j > 0:
                hidden.append(k)
    ch.legend.legendEntry = [LegendEntry(idx=i, delete=True) for i in hidden]
    ws.add_chart(ch, f"A{H0 + 7}")
    ws.column_dimensions["A"].width = 16
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(out)
    print("wrote", out)


# ---------------------------------------------------------------- PowerPoint (native chart)
def build_ppt_chart(rows, deck_in, deck_out, slide_title="System-Level Results 1 – S.2.3"):
    from pptx import Presentation
    from pptx.chart.data import XyChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_MARKER_STYLE
    from pptx.oxml.ns import qn
    from pptx.util import Inches, Pt

    prs = Presentation(deck_in)
    slide = next(s for s in prs.slides
                 if any(sh.has_text_frame and sh.text_frame.text.strip() == slide_title for sh in s.shapes))
    old = [sh for sh in slide.shapes if sh.shape_type == 13 and sh.width > Inches(5)]   # the scatter picture
    x, y, w, h = Inches(2.95), Inches(0.58), Inches(6.85), Inches(4.25)
    for sh in old:
        sh._element.getparent().remove(sh._element)

    cd = XyChartData()
    meta = []                                            # (kind, payload) per series, in order
    s = cd.add_series("Vehicle level (Section 1)", number_format="#,##0")
    for r in rows:
        if on_chart(r):
            s.add_data_point(r["vl_m"], r["vl_c"])
    meta.append(("vl", None))
    for r in rows:
        if on_chart(r):
            s = cd.add_series(f"{r['last']} {r['kind']} (line)")
            s.add_data_point(r["vl_m"], r["vl_c"])
            s.add_data_point(r["sl_m"], r["sl_c"])
            meta.append(("line", r))
    for who in TEAM:
        for j, r in enumerate(legend_order(rows, who)):
            s = cd.add_series(pick_label(r) if r["pick"] else f"{r['s1']} / {r['s2']} ({r['last']})")
            s.add_data_point(r["sl_m"], r["sl_c"])
            meta.append(("sl", (r, j == 0)))

    gf = slide.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER, x, y, w, h, cd)
    chart = gf.chart
    chart.font.name, chart.font.size = "Palatino Linotype", Pt(9)
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.TOP
    chart.legend.include_in_layout = False
    chart.legend.font.size = Pt(8.5)
    for ax, (lo, hi), unit, title, fmt in ((chart.category_axis, XLIM, 200, "Gross liftoff mass (t)", "#,##0"),
                                           (chart.value_axis, YLIM, 0.5, "NRE program cost ($B, FY2025)", "0.0")):
        ax.minimum_scale, ax.maximum_scale, ax.major_unit = lo, hi, unit
        ax.has_major_gridlines = True
        ax.major_gridlines.format.line.color.rgb = RGBColor.from_string("E2E5EA")
        ax.tick_labels.number_format, ax.tick_labels.number_format_is_linked = fmt, False
        ax.tick_labels.font.size = Pt(9)
        ax.has_title = True
        ax.axis_title.text_frame.text = title
        ax.axis_title.text_frame.paragraphs[0].runs[0].font.size = Pt(9.5)
        ax.axis_title.text_frame.paragraphs[0].runs[0].font.bold = False

    def line_alpha(series, rgb, alpha, width_pt, arrow):
        ln = series.format.line
        ln.color.rgb = RGBColor.from_string(rgb)
        ln.width = Pt(width_pt)
        el = series.format.line._get_or_add_ln()
        clr = el.find(qn("a:solidFill")).find(qn("a:srgbClr"))
        etree.SubElement(clr, qn("a:alpha")).set("val", str(alpha * 1000))
        if arrow:
            etree.SubElement(el, qn("a:tailEnd"), type="triangle", w="med", len="med")

    plot = chart.plots[0]
    hidden = []
    for i, (series, (kind, payload)) in enumerate(zip(plot.series, meta)):
        series.smooth = False
        if kind == "vl":
            series.format.line.fill.background()
            series.marker.style, series.marker.size = XL_MARKER_STYLE.CIRCLE, 6
            series.marker.format.fill.solid()
            series.marker.format.fill.fore_color.rgb = RGBColor.from_string(VL_GRAY)
            series.marker.format.line.color.rgb = RGBColor.from_string(VL_GRAY)
        elif kind == "line":
            series.marker.style = XL_MARKER_STYLE.NONE
            line_alpha(series, COLOR[payload["who"]], 65, 1.5, arrow=True)
            hidden.append(i)
        elif kind == "sl":
            r, first = payload
            c = RGBColor.from_string(COLOR[r["who"]])
            series.format.line.fill.background()
            series.marker.style = XL_MARKER_STYLE.CIRCLE if r["kind"] == "min-mass" else XL_MARKER_STYLE.SQUARE
            series.marker.size = 14 if r["pick"] else 10
            series.marker.format.fill.solid()
            series.marker.format.fill.fore_color.rgb = c if r["fit"] else RGBColor.from_string(TINT[r["who"]])
            series.marker.format.line.color.rgb = RGBColor.from_string(INK) if r["pick"] else c
            series.marker.format.line.width = Pt(2)
            if r["pick"]:
                dl = series.points[0].data_label
                dl.has_text_frame = True
                dl.text_frame.text = pick_label(r)
                dl.text_frame.paragraphs[0].runs[0].font.bold = True
                dl.text_frame.paragraphs[0].runs[0].font.size = Pt(9)
                dl.position = XL_LABEL_POSITION.BELOW
            if not first:
                hidden.append(i)
    # drop the per-design lines from the legend
    leg = chart.legend._element
    pos = leg.find(qn("c:legendPos"))
    for k, i in enumerate(hidden):
        e = etree.Element(qn("c:legendEntry"))
        etree.SubElement(e, qn("c:idx")).set("val", str(plot.series[i]._element.find(qn("c:idx")).get("val")))
        etree.SubElement(e, qn("c:delete")).set("val", "1")
        pos.addnext(e) if k == 0 else leg.findall(qn("c:legendEntry"))[-1].addnext(e)

    # the slide's legend text said "ringed" for the old picture
    for sh in slide.shapes:
        if sh.has_text_frame:
            for para in sh.text_frame.paragraphs:
                for run in para.runs:
                    for old, new in (("Ringed = group picks", "Big labeled points = group picks"),
                                     ("hollow = engines don't fit", "light fill = engines don't fit")):
                        run.text = run.text.replace(old, new)
    # note for designs outside the window (the solid first stage)
    note = off_chart_note(rows)
    if note:
        tb = slide.shapes.add_textbox(Inches(2.95), Inches(4.84), Inches(6.85), Inches(0.2))
        tb.text_frame.text = note
        run = tb.text_frame.paragraphs[0].runs[0]
        run.font.size, run.font.italic, run.font.name = Pt(8.5), True, "Palatino Linotype"
        run.font.color.rgb = RGBColor.from_string("4A5568")
    prs.save(deck_out)
    print("wrote", deck_out)


if __name__ == "__main__":
    import os
    here = os.path.dirname(os.path.abspath(__file__))
    report = sys.argv[1] if len(sys.argv) > 1 else os.path.join(here, "csvs", "S2_report.json")
    out = sys.argv[2] if len(sys.argv) > 2 else os.path.join(here, "csvs", "S2_system_level_results.xlsx")
    if not os.path.exists(report):
        sys.exit(f"{report} not found -- run python main_section2.py first")
    rows = load(report)
    build_workbook(rows, out)
    if len(sys.argv) >= 5:
        build_ppt_chart(rows, sys.argv[3], sys.argv[4])
